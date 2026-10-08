"""Conversiones que se hacen dentro de la API (sin Gotenberg).

Cada función recibe la ruta de un PDF y escribe el resultado en `out_path`.
Todo lo que abre un PDF (también contar sus páginas) corre en un proceso aparte
(ver `child_main` y app.py): así se puede cortar si tarda demasiado, tiene un tope
de memoria y, si se cae, no tumba al servidor web.
"""

from __future__ import annotations

import errno
import io
import math
import re

import pymupdf

# Límites para que una sola conversión no acapare el servidor.
MAX_SLIDE_PAGES = 150
SLIDE_DPI = 150
# Tope de píxeles por diapositiva (unos 25 megapíxeles). Una página gigante
# (por ejemplo, un plano de 200 × 200 pulgadas) se dibuja con menos DPI en vez de
# pedir gigas de memoria.
MAX_SLIDE_PIXELS = 25_000_000


def pdf_to_word(pdf_path: str, out_path: str) -> str:
    """PDF a Word con pdf2docx (rehace párrafos, tablas e imágenes)."""
    import logging

    from pdf2docx import Converter

    logging.disable(logging.INFO)  # pdf2docx escribe mucho en el registro
    converter = Converter(pdf_path)
    try:
        converter.convert(out_path, multi_processing=False)
    finally:
        converter.close()
    return "Revisa el documento: los PDF con diseños complejos pueden necesitar algunos ajustes."


def _clean(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").strip()
    # Excel no acepta caracteres de control.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)


def _as_number(text: str):
    """Convierte '1.234,50' o '1,234.50' o '25' en número si se puede."""
    t = text.replace(" ", "").replace(" ", "")
    if not re.fullmatch(r"-?[\d.,]+", t) or not re.search(r"\d", t):
        return text
    if re.fullmatch(r"0\d+", t):
        return text  # códigos con ceros a la izquierda (DNI, cuentas) se dejan como texto
    if "," in t and "." in t:
        if t.rfind(",") > t.rfind("."):
            t = t.replace(".", "").replace(",", ".")
        else:
            t = t.replace(",", "")
    elif "," in t:
        parts = t.split(",")
        t = t.replace(",", ".") if len(parts) == 2 and len(parts[1]) != 3 else t.replace(",", "")
    try:
        number = float(t)
        return int(number) if number.is_integer() and "." not in t else number
    except ValueError:
        return text


def pdf_to_excel(pdf_path: str, out_path: str) -> str:
    """Busca tablas con PyMuPDF y las pasa a una hoja por página."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    doc = pymupdf.open(pdf_path)
    wb = Workbook()
    wb.remove(wb.active)
    tables_found = 0
    try:
        for index, page in enumerate(doc):
            try:
                tables = page.find_tables().tables
            except Exception:  # noqa: BLE001 - una página rara no debe tumbar todo
                tables = []
            if not tables:
                continue
            ws = wb.create_sheet(title=f"Página {index + 1}"[:31])
            row = 1
            for table in tables:
                data = table.extract()
                for r, cells in enumerate(data):
                    for c, value in enumerate(cells):
                        text = _clean(value)
                        cell = ws.cell(row=row, column=c + 1, value=_as_number(text) if text else None)
                        if r == 0:
                            cell.font = Font(bold=True)
                    row += 1
                row += 1  # una fila vacía entre tablas
                tables_found += 1
            for column in ws.columns:
                width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
                ws.column_dimensions[column[0].column_letter].width = min(60, max(8, width + 2))

        if tables_found:
            wb.save(out_path)
            return f"Encontramos {tables_found} {'tabla' if tables_found == 1 else 'tablas'}."

        # Sin tablas: se copia el texto línea por línea para no devolver un archivo vacío.
        ws = wb.create_sheet(title="Texto")
        row = 1
        for index, page in enumerate(doc):
            for line in page.get_text("text").splitlines():
                line = _clean(line)
                if line:
                    ws.cell(row=row, column=1, value=line)
                    row += 1
        if row == 1:
            raise ValueError("sin texto")
        ws.column_dimensions["A"].width = 100
        wb.save(out_path)
        return "No encontramos tablas, así que copiamos el texto línea por línea."
    finally:
        doc.close()


def slide_zoom(width_pt: float, height_pt: float) -> float:
    """Aumento para dibujar una página: SLIDE_DPI, o menos si pasaría de MAX_SLIDE_PIXELS.

    1 punto = 1/72 de pulgada, así que a `zoom` la imagen mide (ancho × zoom) × (alto × zoom) píxeles.
    """
    zoom = SLIDE_DPI / 72
    area = max(width_pt * height_pt, 1.0)
    if area * zoom * zoom > MAX_SLIDE_PIXELS:
        zoom = math.sqrt(MAX_SLIDE_PIXELS / area)
    return zoom


def pdf_to_powerpoint(pdf_path: str, out_path: str) -> str:
    """Una diapositiva por página, con la página como imagen y su texto en las notas."""
    from pptx import Presentation
    from pptx.util import Emu

    doc = pymupdf.open(pdf_path)
    try:
        if doc.page_count > MAX_SLIDE_PAGES:
            raise ValueError(f"demasiadas páginas ({doc.page_count})")
        prs = Presentation()
        first = doc[0].rect
        # 1 punto = 12 700 EMU. PowerPoint acepta de 1 a 56 pulgadas por lado.
        scale = min(1.0, 56 * 72 / max(first.width, first.height))
        prs.slide_width = Emu(int(first.width * scale * 12700))
        prs.slide_height = Emu(int(first.height * scale * 12700))
        blank = prs.slide_layouts[6]
        for page in doc:
            zoom = slide_zoom(page.rect.width, page.rect.height)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            image = io.BytesIO(pix.tobytes("png"))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(image, 0, 0, width=prs.slide_width, height=prs.slide_height)
            text = page.get_text("text").strip()
            if text:
                slide.notes_slide.notes_text_frame.text = text[:5000]
        prs.save(out_path)
        return "Cada página quedó como imagen en su diapositiva, con su texto en las notas."
    finally:
        doc.close()


def page_count(pdf_path: str) -> int:
    with pymupdf.open(pdf_path) as doc:
        if doc.needs_pass:
            raise PermissionError("protegido")
        return doc.page_count


LOCAL_CONVERTERS = {
    "pdf-a-word": (pdf_to_word, "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf-a-excel": (pdf_to_excel, "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf-a-powerpoint": (
        pdf_to_powerpoint,
        "pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ),
}


COUNT_PAGES = "contar-paginas"


def _limit_child(memory_bytes: int, file_bytes: int) -> None:
    """Topes del proceso hijo, para que un PDF raro no se coma el servidor."""
    import resource

    try:
        # Si aun así falta memoria, que el sistema mate primero a este hijo y no al servidor web.
        with open("/proc/self/oom_score_adj", "w") as fh:
            fh.write("1000")
    except OSError:
        pass
    # RLIMIT_AS: memoria máxima. Si se pasa, la conversión falla con MemoryError.
    # RLIMIT_FSIZE: tamaño máximo de cada archivo que escribe (el resultado).
    for limit, value in ((resource.RLIMIT_AS, memory_bytes), (resource.RLIMIT_FSIZE, file_bytes)):
        if value > 0:
            try:
                resource.setrlimit(limit, (value, value))
            except (ValueError, OSError):
                pass


def _reason(exc: BaseException) -> str:
    """Pone nombre a los errores que la API sabe explicar."""
    text = str(exc)
    if isinstance(exc, MemoryError) or "malloc" in text or "out of memory" in text.lower():
        return "memoria"
    if isinstance(exc, OSError) and exc.errno == errno.EFBIG:
        return "muy-grande"
    if isinstance(exc, PermissionError) or text == "protegido":
        return "protegido"
    if "demasiadas páginas" in text:
        return "demasiadas-paginas"
    if "sin texto" in text:
        return "sin-texto"
    return "otro"


def child_main(task: str, args: tuple, conn, memory_bytes: int = 0, file_bytes: int = 0) -> None:
    """Punto de entrada del proceso hijo: hace la tarea y avisa el resultado por la tubería.

    Envía ("ok", resultado) o ("error", motivo, detalle).
    """
    try:
        _limit_child(memory_bytes, file_bytes)
        if task == COUNT_PAGES:
            result = page_count(*args)
        else:
            result = LOCAL_CONVERTERS[task][0](*args) or ""
        conn.send(("ok", result))
    except BaseException as exc:  # noqa: BLE001 - cualquier error se avisa a la API
        conn.send(("error", _reason(exc), f"{type(exc).__name__}: {exc}"[:300]))
    finally:
        conn.close()
