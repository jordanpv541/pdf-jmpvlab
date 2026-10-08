"""Pruebas de la API sin Gotenberg real (se reemplaza por uno falso con httpx.MockTransport).

    cd server && python -m pytest tests -q

Las pruebas con Gotenberg de verdad están en tests/integration/ y corren en GitHub Actions.
"""

import asyncio
import io
import socket
import struct
import sys
import tempfile
import threading
import time
import zipfile
from pathlib import Path
from urllib.parse import unquote

import httpx
import pymupdf
import pytest
import python_multipart
import uvicorn
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "api"))

import app as api  # noqa: E402
import converters  # noqa: E402

import child_helpers  # noqa: E402

ORIGIN = "https://pdf.jmpvlab.com"


def make_pdf(pages=2, table=False, password=None, size=None) -> bytes:
    doc = pymupdf.open()
    for n in range(pages):
        page = doc.new_page(**({"width": size[0], "height": size[1]} if size else {}))
        page.insert_text((72, 72), f"Página de prueba {n + 1}", fontsize=18)
        if table:
            x0, y0, w, h = 72, 120, 120, 24
            rows = [["Producto", "Cantidad", "Precio"], ["Lápiz", "12", "1,50"], ["Cuaderno", "3", "1.234,90"], ["Código", "00123", "7"]]
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    rect = pymupdf.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
                    page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                    page.insert_text((rect.x0 + 4, rect.y1 - 7), value, fontsize=11)
    kwargs = {}
    if password:
        kwargs = {"encryption": pymupdf.PDF_ENCRYPT_AES_256, "user_pw": password, "owner_pw": password + "x"}
    data = doc.tobytes(**kwargs)
    doc.close()
    return data


@pytest.fixture(autouse=True)
def reset_limits(monkeypatch, tmp_path):
    api._hits.clear()
    api._in_flight_by_ip.clear()
    monkeypatch.setattr(api, "_in_flight", 0)
    monkeypatch.setattr(api, "_slots", asyncio.Semaphore(api.MAX_JOBS))
    monkeypatch.setattr(api, "RATE_LIMIT", 30)
    # Las carpetas de cada conversión se crean aquí, para revisar que se borren.
    jobs = tmp_path / "jobs"
    jobs.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(jobs))
    yield
    assert not list(jobs.glob("conv-*")), "quedó una carpeta temporal sin borrar"


@pytest.fixture
def client():
    return TestClient(api.app)


def parse_multipart(request: httpx.Request):
    fields, files = {}, []

    def on_field(field):
        fields[field.field_name.decode()] = (field.value or b"").decode()

    def on_file(file):
        files.append(file.file_name.decode())

    python_multipart.parse_form({"Content-Type": request.headers["content-type"]}, io.BytesIO(request.content), on_field, on_file)
    return fields, files


class FakeGotenberg(list):
    """Lista de lo que recibió Gotenberg; `reply` cambia lo que responde."""

    reply: dict


@pytest.fixture
def fake_gotenberg(monkeypatch):
    """Gotenberg falso: anota lo que recibe y responde un PDF (o lo que diga `reply`)."""
    calls = FakeGotenberg()
    calls.reply = {"status": 200, "content": None, "stream": None}

    def handler(request: httpx.Request):
        assert request.headers["content-type"].startswith("multipart/form-data"), "Gotenberg solo acepta multipart"
        fields, files = parse_multipart(request)
        calls.append((request.url.path, files, fields))
        reply = calls.reply
        if reply["stream"] is not None:
            return httpx.Response(reply["status"], content=reply["stream"]())
        return httpx.Response(reply["status"], content=reply["content"] or make_pdf(1))

    monkeypatch.setattr(api, "_gotenberg_transport", httpx.MockTransport(handler))
    return calls


def post(client, kind, name, content, headers=None, **fields):
    return client.post(
        f"/v1/convert/{kind}",
        files={"file": (name, content, "application/octet-stream")},
        data=fields,
        headers={"Origin": ORIGIN, **(headers or {})},
    )


# ---------- Lo básico ----------


def test_health(client):
    r = client.get("/v1/health")
    assert r.status_code == 200 and r.json()["ok"] is True
    assert "github.com" in r.json()["source"]
    assert r.headers["x-content-type-options"] == "nosniff"


def test_cors_only_for_our_site(client):
    ok = client.options("/v1/convert/pdf-a-word", headers={"Origin": ORIGIN, "Access-Control-Request-Method": "POST"})
    bad = client.options("/v1/convert/pdf-a-word", headers={"Origin": "https://otro.com", "Access-Control-Request-Method": "POST"})
    assert ok.headers.get("access-control-allow-origin") == ORIGIN
    assert bad.headers.get("access-control-allow-origin") is None


def test_unknown_kind(client):
    assert post(client, "nada", "a.pdf", make_pdf()).status_code == 404


def test_word_to_pdf_uses_libreoffice(client, fake_gotenberg):
    docx = io.BytesIO()
    with zipfile.ZipFile(docx, "w") as z:
        z.writestr("word/document.xml", "<x/>")
    r = post(client, "word-a-pdf", "Mi informe.docx", docx.getvalue())
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert "Mi%20informe.pdf" in r.headers["content-disposition"]
    assert r.content.startswith(b"%PDF")
    assert fake_gotenberg[0][0] == "/forms/libreoffice/convert"
    assert fake_gotenberg[0][1] == ["documento.docx"]


def test_rejects_wrong_extension_and_fake_content(client, fake_gotenberg):
    assert post(client, "word-a-pdf", "foto.jpg", b"\xff\xd8\xff").status_code == 415
    r = post(client, "word-a-pdf", "falso.docx", b"esto no es un docx")
    assert r.status_code == 415 and "no parece" in r.json()["error"]
    assert not fake_gotenberg


def test_empty_file(client, fake_gotenberg):
    assert post(client, "pdf-a-word", "vacio.pdf", b"").status_code == 400


def test_too_big(client, fake_gotenberg, monkeypatch):
    monkeypatch.setattr(api, "MAX_BYTES", 1000)
    r = post(client, "pdf-a-pdfa", "grande.pdf", b"%PDF-1.7\n" + b"0" * 5000)
    assert r.status_code == 413
    assert not fake_gotenberg


def test_pdfa_level(client, fake_gotenberg):
    r = post(client, "pdf-a-pdfa", "a.pdf", make_pdf(), pdfa="PDF/A-1b")
    assert r.status_code == 200
    assert fake_gotenberg[0][2]["pdfa"] == "PDF/A-1b"
    post(client, "pdf-a-pdfa", "a.pdf", make_pdf(), pdfa="cualquier cosa")
    assert fake_gotenberg[1][2]["pdfa"] == "PDF/A-2b"


def test_password_pdf(client):
    r = post(client, "pdf-a-word", "clave.pdf", make_pdf(password="abc"))
    assert r.status_code == 422 and "contraseña" in r.json()["error"]


def test_damaged_pdf(client):
    r = post(client, "pdf-a-word", "roto.pdf", b"%PDF-1.7\nesto no es un pdf de verdad")
    assert r.status_code == 422 and "Reparar PDF" in r.json()["error"]


def test_pdf_to_word(client):
    r = post(client, "pdf-a-word", "doc.pdf", make_pdf(2))
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"].endswith("doc.docx")
    from docx import Document

    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "Página de prueba 1" in text


def test_pdf_to_excel_tables(client):
    r = post(client, "pdf-a-excel", "tabla.pdf", make_pdf(1, table=True))
    assert r.status_code == 200, r.text
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    flat = [v for row in values for v in row if v is not None]
    assert "Producto" in flat and "Lápiz" in flat
    assert 1.5 in flat and 1234.9 in flat  # números en formato español
    assert "00123" in flat  # los códigos con ceros se quedan como texto
    assert "1 tabla" in unquote(r.headers["x-result-note"])


def test_pdf_to_excel_without_tables(client):
    r = post(client, "pdf-a-excel", "texto.pdf", make_pdf(2))
    assert r.status_code == 200
    assert "No encontramos tablas" in unquote(r.headers["x-result-note"])


def test_pdf_to_powerpoint(client):
    r = post(client, "pdf-a-powerpoint", "presentacion.pdf", make_pdf(3))
    assert r.status_code == 200, r.text
    from pptx import Presentation

    prs = Presentation(io.BytesIO(r.content))
    assert len(prs.slides) == 3
    assert "Página de prueba 2" in prs.slides[1].notes_slide.notes_text_frame.text


def test_page_limit(client, monkeypatch):
    monkeypatch.setattr(api, "MAX_PDF_PAGES", 2)
    r = post(client, "pdf-a-word", "largo.pdf", make_pdf(3))
    assert r.status_code == 413 and "3 páginas" in r.json()["error"]


def test_word_has_lower_page_limit(client, fake_gotenberg, monkeypatch):
    monkeypatch.setattr(api, "MAX_WORD_PAGES", 2)
    r = post(client, "pdf-a-word", "largo.pdf", make_pdf(3))
    assert r.status_code == 413
    assert "Word" in r.json()["error"] and "Dividir PDF" in r.json()["error"]
    # Las otras conversiones siguen con el límite general.
    assert post(client, "pdf-a-pdfa", "largo.pdf", make_pdf(3)).status_code == 200


def test_rate_limit(client, fake_gotenberg, monkeypatch):
    monkeypatch.setattr(api, "RATE_LIMIT", 2)
    for _ in range(2):
        assert post(client, "pdf-a-pdfa", "a.pdf", make_pdf()).status_code == 200
    r = post(client, "pdf-a-pdfa", "a.pdf", make_pdf())
    assert r.status_code == 429 and "Vuelve a intentarlo" in r.json()["error"]
    assert r.headers.get("access-control-allow-origin") == ORIGIN  # la web puede leer el mensaje


def test_html_file(client, fake_gotenberg):
    r = post(client, "html-a-pdf", "pagina.html", b"<h1>Hola</h1>", paper="carta")
    assert r.status_code == 200
    route, names, data = fake_gotenberg[0]
    assert route == "/forms/chromium/convert/html" and names == ["index.html"]
    assert data["paperWidth"] == "8.5"


def test_form_is_not_parsed_by_fastapi(client, fake_gotenberg, monkeypatch):
    """El archivo se guarda una sola vez, directo en la carpeta de la conversión."""

    async def no_form(*_args, **_kwargs):
        raise AssertionError("no debe usarse request.form(): guarda una copia extra del archivo")

    monkeypatch.setattr("starlette.requests.Request.form", no_form)
    assert post(client, "pdf-a-pdfa", "a.pdf", make_pdf()).status_code == 200


# ---------- Direcciones web ----------


@pytest.mark.parametrize(
    "url",
    [
        "http://127.0.0.1/",
        "http://localhost:3000/",
        "http://169.254.169.254/latest/meta-data/",
        "http://10.0.0.5/",
        "http://192.168.1.1/admin",
        "http://100.64.0.1/",
        "http://[::1]/",
        "http://[::ffff:127.0.0.1]/",
        "http://[fd00::1]/",
        "file:///etc/passwd",
        "ftp://example.com/",
        "https://user:pw@example.com/",
        "https://example.com:8443/",
        "http://gotenberg.internal/",
        "http://api:8000/v1/health",
        "http://api/v1/health",
        "http://gotenberg/health",
        "http://2130706433/",
    ],
)
def test_html_url_blocks_internal_addresses(client, fake_gotenberg, url):
    r = client.post("/v1/convert/html-a-pdf", data={"url": url}, headers={"Origin": ORIGIN})
    assert r.status_code == 400, (url, r.text)
    assert not fake_gotenberg


@pytest.mark.parametrize("address", ["127.0.0.1", "10.1.2.3", "100.64.5.5", "169.254.169.254", "::1", "fd12::1", "::ffff:10.0.0.1"])
def test_html_url_name_that_points_inside(client, fake_gotenberg, monkeypatch, address):
    """Un nombre público (como 127.0.0.1.nip.io) que apunta a una IP interna también se rechaza."""

    async def fake_resolve(host, port):
        return ["93.184.215.14", address]

    monkeypatch.setattr(api, "resolve", fake_resolve)
    r = client.post("/v1/convert/html-a-pdf", data={"url": "https://127.0.0.1.nip.io/"}, headers={"Origin": ORIGIN})
    assert r.status_code == 400 and "pública" in r.json()["error"]
    assert not fake_gotenberg


def test_html_url_public_ip(client, fake_gotenberg):
    r = client.post("/v1/convert/html-a-pdf", data={"url": "http://1.1.1.1/"}, headers={"Origin": ORIGIN})
    assert r.status_code == 200, r.text
    route, files, fields = fake_gotenberg[0]
    assert route == "/forms/chromium/convert/url" and files == []
    assert fields["url"] == "http://1.1.1.1/"  # llegó como multipart, que es lo que Gotenberg exige


def test_html_url_public_ipv6(client, fake_gotenberg):
    r = client.post("/v1/convert/html-a-pdf", files={"url": (None, "http://[2606:4700:4700::1111]/")}, headers={"Origin": ORIGIN})
    assert r.status_code == 200, r.text


def test_url_page_with_error_status(client, fake_gotenberg):
    """Gotenberg responde 409 si la página contestó 404/500: no es un «archivo dañado»."""
    fake_gotenberg.reply["status"] = 409
    fake_gotenberg.reply["content"] = b"Invalid HTTP status code from the main page: 404"
    r = client.post("/v1/convert/html-a-pdf", data={"url": "http://1.1.1.1/no-existe"}, headers={"Origin": ORIGIN})
    assert r.status_code == 422 and "La página no cargó" in r.json()["error"]

    fake_gotenberg.reply["status"] = 403
    r = client.post("/v1/convert/html-a-pdf", data={"url": "http://1.1.1.1/"}, headers={"Origin": ORIGIN})
    assert r.status_code == 400 and "no es pública" in r.json()["error"]

    fake_gotenberg.reply["status"] = 409
    r = post(client, "html-a-pdf", "pagina.html", b"<h1>Hola</h1>")
    assert r.status_code == 422 and "No se pudo convertir el archivo" in r.json()["error"]


def test_gotenberg_busy_and_down(client, fake_gotenberg, monkeypatch):
    fake_gotenberg.reply["status"] = 503
    assert post(client, "pdf-a-pdfa", "a.pdf", make_pdf()).status_code == 503

    def down(request):
        raise httpx.ConnectError("sin conexión")

    monkeypatch.setattr(api, "_gotenberg_transport", httpx.MockTransport(down))
    r = post(client, "pdf-a-pdfa", "a.pdf", make_pdf())
    assert r.status_code == 503 and "no está disponible" in r.json()["error"]


# ---------- Tope del archivo convertido ----------


def test_output_size_cap_declared(client, fake_gotenberg, monkeypatch):
    monkeypatch.setattr(api, "MAX_OUTPUT_BYTES", 1000)
    fake_gotenberg.reply["content"] = b"%PDF-1.7\n" + b"0" * 5000
    r = post(client, "pdf-a-pdfa", "a.pdf", make_pdf())
    assert r.status_code == 422 and "pesa más de" in r.json()["error"]


def test_output_size_cap_streamed(client, fake_gotenberg, monkeypatch):
    """Sin Content-Length: se corta mientras llega, sin guardarlo todo en memoria."""
    monkeypatch.setattr(api, "MAX_OUTPUT_BYTES", 1000)
    sent = []

    async def endless():
        for _ in range(1000):
            sent.append(1)
            yield b"0" * 500

    fake_gotenberg.reply["stream"] = endless
    r = post(client, "pdf-a-pdfa", "a.pdf", make_pdf())
    assert r.status_code == 422 and "pesa más de" in r.json()["error"]
    assert len(sent) < 10


def test_output_streamed_from_disk(client, fake_gotenberg):
    big = make_pdf(1) + b"%" + b"x" * 300_000
    fake_gotenberg.reply["content"] = big
    r = post(client, "pdf-a-pdfa", "a.pdf", make_pdf())
    assert r.status_code == 200 and r.content == big


# ---------- Límites antes de leer el archivo ----------


def call_asgi(headers, path="/v1/convert/pdf-a-word", client=("203.0.113.7", 5000)):
    """Llama a la API directo (ASGI) y anota si alguien intentó leer el cuerpo."""
    reads = []
    sent = []

    async def receive():
        reads.append(1)
        await asyncio.sleep(3600)  # el cuerpo nunca llega

    async def send(message):
        sent.append(message)

    scope = {
        "type": "http",
        "asgi": {"version": "3.0"},
        "http_version": "1.1",
        "method": "POST",
        "scheme": "http",
        "path": path,
        "raw_path": path.encode(),
        "query_string": b"",
        "root_path": "",
        "headers": [(b"host", b"api"), (b"origin", ORIGIN.encode())] + headers,
        "client": client,
        "server": ("api", 8000),
        "state": {},
    }
    asyncio.run(asyncio.wait_for(api.app(scope, receive, send), 5))
    start = next(m for m in sent if m["type"] == "http.response.start")
    body = b"".join(m.get("body", b"") for m in sent if m["type"] == "http.response.body")
    return start["status"], dict(start["headers"]), body, reads


def test_early_reject_too_big_without_reading(monkeypatch):
    status, headers, body, reads = call_asgi([(b"content-type", b"multipart/form-data; boundary=x"), (b"content-length", str(api.MAX_BYTES * 2).encode())])
    assert status == 413 and "MB" in body.decode()
    assert headers[b"access-control-allow-origin"] == ORIGIN.encode()
    assert reads == []  # el cuerpo no se tocó


def test_early_reject_without_length():
    status, _, _, reads = call_asgi([(b"content-type", b"multipart/form-data; boundary=x"), (b"transfer-encoding", b"chunked")])
    assert status == 411 and reads == []
    status, _, _, reads = call_asgi([(b"content-type", b"multipart/form-data; boundary=x")])
    assert status == 411 and reads == []


def test_early_reject_unknown_kind():
    status, _, _, reads = call_asgi([(b"content-length", b"10")], path="/v1/convert/nada")
    assert status == 404 and reads == []


def test_per_ip_limit_before_reading(monkeypatch):
    api._in_flight_by_ip["203.0.113.7"] = api.MAX_PER_IP
    status, _, body, reads = call_asgi([(b"content-length", b"1000")])
    assert status == 429 and "Ya estás convirtiendo" in body.decode() and reads == []
    # Otra persona sí puede.
    api._in_flight_by_ip.clear()
    monkeypatch.setattr(api, "_in_flight", api.MAX_JOBS + api.MAX_WAITING)
    status, _, body, reads = call_asgi([(b"content-length", b"1000")], client=("198.51.100.1", 1))
    assert status == 503 and "mucha gente" in body.decode() and reads == []


def test_rate_limit_before_reading(monkeypatch):
    monkeypatch.setattr(api, "RATE_LIMIT", 0)
    status, _, _, reads = call_asgi([(b"content-length", b"1000")])
    assert status == 429 and reads == []


def test_per_ip_key_is_forwarded_ip(client, fake_gotenberg):
    """Detrás de Caddy, la IP real viene en X-Forwarded-For (Caddy borra la que mande el visitante)."""
    api._in_flight_by_ip["198.51.100.9"] = api.MAX_PER_IP
    r = post(client, "pdf-a-pdfa", "a.pdf", make_pdf(), headers={"X-Forwarded-For": "198.51.100.9"})
    assert r.status_code == 429
    r = post(client, "pdf-a-pdfa", "a.pdf", make_pdf(), headers={"X-Forwarded-For": "198.51.100.10"})
    assert r.status_code == 200
    assert api._in_flight == 0 and "198.51.100.10" not in api._in_flight_by_ip  # se liberó al terminar


# ---------- Con un servidor de verdad (uvicorn) y conexiones lentas ----------


@pytest.fixture
def live_server():
    config = uvicorn.Config(api.app, host="127.0.0.1", port=0, http="h11", log_level="warning", lifespan="off")
    server = uvicorn.Server(config)
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    for _ in range(100):
        if server.started:
            break
        time.sleep(0.05)
    port = server.servers[0].sockets[0].getsockname()[1]
    yield port
    server.should_exit = True
    thread.join(10)


def raw_request(port, length, body=b"", chunked=False):
    sock = socket.create_connection(("127.0.0.1", port), timeout=5)
    head = "POST /v1/convert/pdf-a-word HTTP/1.1\r\nHost: api\r\nContent-Type: multipart/form-data; boundary=XyZ\r\n"
    head += "Transfer-Encoding: chunked\r\n" if chunked else f"Content-Length: {length}\r\n"
    sock.sendall(head.encode() + b"\r\n" + body)
    return sock


def read_status(sock, timeout=5) -> int:
    """Lee la respuesta completa y comprueba que el servidor cierre la conexión (no se queda leyendo el resto)."""
    sock.settimeout(timeout)
    data = b""
    while chunk := sock.recv(4096):
        data += chunk
    assert b"connection: close" in data.lower()
    return int(data.split(b" ", 2)[1])


def wait_for(condition, seconds=5):
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        if condition():
            return True
        time.sleep(0.05)
    return False


STALLED_BODY = b'--XyZ\r\nContent-Disposition: form-data; name="file"; filename="a.pdf"\r\nContent-Type: application/pdf\r\n\r\n%PDF-1.7\n'


def test_live_rejects_before_body_arrives(live_server):
    port = live_server
    # Dos subidas que se quedan a medias: ocupan los 2 cupos de esta IP.
    stalled = [raw_request(port, 10_000_000, STALLED_BODY) for _ in range(api.MAX_PER_IP)]
    assert wait_for(lambda: api._in_flight_by_ip.get("127.0.0.1") == api.MAX_PER_IP)

    # La tercera recibe 429 al instante, aunque no mandó ni un byte del archivo.
    third = raw_request(port, 10_000_000)
    assert read_status(third) == 429
    # Demasiado grande: 413 sin mandar el cuerpo.
    big = raw_request(port, 500_000_000)
    assert read_status(big) == 413
    # Sin Content-Length (chunked): 411.
    chunked = raw_request(port, 0, chunked=True)
    assert read_status(chunked) == 411

    # Al cortarse las subidas, los cupos se liberan y las carpetas se borran.
    for sock in stalled + [third, big, chunked]:
        sock.close()
    assert wait_for(lambda: api._in_flight == 0 and not api._in_flight_by_ip)


def test_live_slow_upload_times_out(live_server, monkeypatch):
    monkeypatch.setattr(api, "UPLOAD_TIMEOUT", 1)
    sock = raw_request(live_server, 10_000_000, STALLED_BODY)
    assert read_status(sock, timeout=10) == 408
    sock.close()
    assert wait_for(lambda: api._in_flight == 0)


# ---------- Procesos hijos: tiempo, memoria y caídas ----------


def test_page_count_timeout_does_not_block(monkeypatch, tmp_path):
    """Un PDF que cuelga al abrirse se corta, y mientras tanto la API sigue respondiendo."""
    monkeypatch.setattr(converters, "child_main", child_helpers.hang)
    monkeypatch.setattr(api, "PAGE_COUNT_TIMEOUT", 1)
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(make_pdf())

    async def main():
        ticks = 0

        async def ticker():
            nonlocal ticks
            while True:
                await asyncio.sleep(0.05)
                ticks += 1

        task = asyncio.create_task(ticker())
        started = time.monotonic()
        with pytest.raises(api.Problem) as info:
            await api.count_pages(pdf)
        task.cancel()
        return info.value, time.monotonic() - started, ticks

    problem, elapsed, ticks = asyncio.run(main())
    assert problem.status_code == 422 and "Reparar PDF" in problem.detail
    assert elapsed < 10
    assert ticks >= 10  # el bucle de eventos siguió libre


def test_child_crash_is_reported(monkeypatch, tmp_path):
    monkeypatch.setattr(converters, "child_main", child_helpers.crash)
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(make_pdf())
    result = asyncio.run(api.run_child("pdf-a-word", (str(pdf), str(tmp_path / "x.docx")), 30))
    assert result[0] == "crash"
    with pytest.raises(api.Problem) as info:
        asyncio.run(api.run_local("pdf-a-word", pdf, tmp_path / "x.docx"))
    assert info.value.status_code == 500


def test_child_memory_cap(monkeypatch, tmp_path):
    """Con poca memoria permitida, el hijo falla (y no el servidor)."""
    monkeypatch.setattr(api, "CHILD_MEMORY_MB", 120)
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(make_pdf())
    result = asyncio.run(api.run_child("pdf-a-word", (str(pdf), str(tmp_path / "x.docx")), 60))
    assert result[0] in ("error", "crash"), result
    # Con la memoria normal, la misma conversión funciona.
    monkeypatch.setattr(api, "CHILD_MEMORY_MB", 1536)
    result = asyncio.run(api.run_child("pdf-a-word", (str(pdf), str(tmp_path / "x.docx")), 60))
    assert result[0] == "ok", result


def test_slide_zoom_caps_pixels():
    a4 = converters.slide_zoom(595, 842)
    assert a4 == pytest.approx(150 / 72)  # una página normal queda a 150 DPI
    huge = 200 * 72  # 200 × 200 pulgadas
    zoom = converters.slide_zoom(huge, huge)
    assert (huge * zoom) ** 2 <= converters.MAX_SLIDE_PIXELS
    assert (huge * zoom) ** 2 > converters.MAX_SLIDE_PIXELS * 0.99
    wide = converters.slide_zoom(14400, 10)  # muy angosta: no hace falta bajar
    assert wide == pytest.approx(150 / 72)


def png_size(blob: bytes) -> tuple[int, int]:
    return struct.unpack(">II", blob[16:24])


def test_huge_page_to_powerpoint(client):
    r = post(client, "pdf-a-powerpoint", "plano.pdf", make_pdf(1, size=(14400, 14400)))
    assert r.status_code == 200, r.text
    from pptx import Presentation

    picture = Presentation(io.BytesIO(r.content)).slides[0].shapes[0]
    width, height = png_size(picture.image.blob)
    assert width * height <= converters.MAX_SLIDE_PIXELS * 1.01
