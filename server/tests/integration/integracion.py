"""Pruebas de integración con Gotenberg, Caddy y la API de verdad (las corre GitHub Actions).

Antes hay que levantar los contenedores (desde la carpeta server/):

    cp .env.example .env
    sed -i 's|^API_DOMAIN=.*|API_DOMAIN=localhost|; s|^RATE_LIMIT=.*|RATE_LIMIT=1000|' .env
    docker compose -f docker-compose.yml -f tests/integration/compose.ci.yml up -d --build --wait
    python tests/integration/integracion.py

El archivo no se llama test_*.py a propósito: `pytest server/tests` no lo ejecuta.
Necesita las librerías de api/requirements.txt (httpx, PyMuPDF, python-docx, openpyxl, python-pptx).
"""

from __future__ import annotations

import io
import json
import os
import socket
import ssl
import subprocess
import sys
import threading
import time
import traceback
from pathlib import Path
from urllib.parse import urlparse

import httpx
import pymupdf

SERVER_DIR = Path(__file__).resolve().parents[2]
API = os.getenv("API_URL", "http://127.0.0.1:8000")  # la API directa (compose.ci.yml la publica solo en este equipo)
CADDY = os.getenv("CADDY_URL", "https://localhost")  # la API a través de Caddy, como en el servidor
COMPOSE = ["docker", "compose", "-f", "docker-compose.yml", "-f", "tests/integration/compose.ci.yml"]
ORIGIN = {"Origin": "https://pdf.jmpvlab.com"}

CHECKS = []


def check(fn):
    CHECKS.append(fn)
    return fn


def run(names: list[str]) -> int:
    """Corre las pruebas (todas, o solo las nombradas) sin cortar al primer fallo."""
    results = []
    for fn in CHECKS:
        if names and fn.__name__ not in names:
            continue
        started = time.monotonic()
        try:
            fn()
            results.append((fn.__name__, ""))
            print(f"OK     {fn.__name__} ({time.monotonic() - started:.1f} s)", flush=True)
        except Exception as exc:  # noqa: BLE001
            results.append((fn.__name__, f"{type(exc).__name__}: {exc}"))
            print(f"FALLA  {fn.__name__}: {exc}", flush=True)
            traceback.print_exc()
    failed = [(name, error) for name, error in results if error]
    print(f"\n{len(results) - len(failed)} de {len(results)} pruebas pasaron.")
    for name, error in failed:
        print(f"  FALLA {name}: {error}")
    return 1 if failed or not results else 0


def convert(kind: str, filename: str | None = None, content: bytes | None = None, base: str = API, **fields) -> httpx.Response:
    files = {"file": (filename, content, "application/octet-stream")} if filename else {}
    files.update({name: (None, value) for name, value in fields.items()})  # como la web: siempre multipart
    return httpx.post(f"{base}/v1/convert/{kind}", files=files, headers=ORIGIN, timeout=240, verify=False)


def pdf_text(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def expect_pdf_with(response: httpx.Response, text: str) -> str:
    assert response.status_code == 200, f"{response.status_code} {response.text[:300]}"
    assert response.content.startswith(b"%PDF"), "la respuesta no es un PDF"
    content = pdf_text(response.content)
    assert text in content, f"el PDF no dice «{text}»: {content[:300]!r}"
    return content


def sample_pdf(pages: int = 2) -> bytes:
    doc = pymupdf.open()
    for n in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Pagina de prueba {n + 1}", fontsize=18)
        rows = [["Producto", "Cantidad"], ["Lapiz", "12"], ["Cuaderno", "3"]]
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                rect = pymupdf.Rect(72 + c * 150, 120 + r * 24, 72 + (c + 1) * 150, 144 + r * 24)
                page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                page.insert_text((rect.x0 + 4, rect.y1 - 7), value, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


# ---------- Que todo responda ----------


@check
def salud():
    assert httpx.get(f"{API}/v1/health", timeout=10).json()["ok"] is True
    r = httpx.get(f"{CADDY}/v1/health", timeout=10, verify=False)
    assert r.status_code == 200 and r.json()["ok"] is True
    assert r.headers.get("x-content-type-options") == "nosniff"


# ---------- Conversiones con Gotenberg (LibreOffice y Chromium) ----------


@check
def word_a_pdf():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hola Word jmpvlab")
    out = io.BytesIO()
    doc.save(out)
    expect_pdf_with(convert("word-a-pdf", "informe.docx", out.getvalue()), "Hola Word jmpvlab")


@check
def excel_a_pdf():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "Hola Excel jmpvlab"
    out = io.BytesIO()
    wb.save(out)
    expect_pdf_with(convert("excel-a-pdf", "tabla.xlsx", out.getvalue()), "Hola Excel jmpvlab")


@check
def powerpoint_a_pdf():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Hola PowerPoint jmpvlab"
    out = io.BytesIO()
    prs.save(out)
    expect_pdf_with(convert("powerpoint-a-pdf", "charla.pptx", out.getvalue()), "Hola PowerPoint jmpvlab")


@check
def html_a_pdf():
    html = b"<!doctype html><meta charset=utf-8><h1>Hola HTML jmpvlab</h1>"
    expect_pdf_with(convert("html-a-pdf", "pagina.html", html, paper="carta"), "Hola HTML jmpvlab")


@check
def direccion_a_pdf():
    last = None
    for _ in range(3):  # example.com es externo: se reintenta si falla la red
        last = convert("html-a-pdf", url="https://example.com/")
        if last.status_code == 200:
            break
        time.sleep(5)
    # El título de example.com no siempre sale al extraer el texto del PDF: se busca el párrafo.
    expect_pdf_with(last, "documentation examples")


@check
def direccion_con_error_404():
    """Gotenberg responde 409 si la página contesta 404: la API lo explica bien."""
    r = convert("html-a-pdf", url="https://github.com/jordanpv541/pdf-jmpvlab/esta-pagina-no-existe")
    assert r.status_code == 422, f"{r.status_code} {r.text[:200]}"
    assert "La página no cargó" in r.json()["error"], r.text


# ---------- Conversiones que parten de un PDF ----------


@check
def pdf_a_word():
    from docx import Document

    r = convert("pdf-a-word", "doc.pdf", sample_pdf())
    assert r.status_code == 200, r.text[:300]
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "Pagina de prueba 1" in text, text[:300]


@check
def pdf_a_excel():
    from openpyxl import load_workbook

    r = convert("pdf-a-excel", "tabla.pdf", sample_pdf(1))
    assert r.status_code == 200, r.text[:300]
    values = [c.value for ws in load_workbook(io.BytesIO(r.content)).worksheets for row in ws.iter_rows() for c in row]
    assert "Producto" in values, values[:20]


@check
def pdf_a_powerpoint():
    from pptx import Presentation

    r = convert("pdf-a-powerpoint", "charla.pdf", sample_pdf(3))
    assert r.status_code == 200, r.text[:300]
    assert len(Presentation(io.BytesIO(r.content)).slides) == 3


@check
def pdf_a_pdfa():
    r = convert("pdf-a-pdfa", "doc.pdf", sample_pdf(), pdfa="PDF/A-2b")
    assert r.status_code == 200, r.text[:300]
    with pymupdf.open(stream=r.content, filetype="pdf") as doc:
        xmp = doc.get_xml_metadata()
    assert "pdfaid" in xmp, f"sin metadatos PDF/A: {xmp[:300]!r}"


# ---------- SSRF: nada de direcciones internas ----------

INTERNAL_URLS = [
    "http://169.254.169.254/latest/meta-data/",
    "http://127.0.0.1/",
    "http://api:8000/v1/health",
    "http://gotenberg:3000/health",
    "http://[::1]/",
    "http://127.0.0.1.nip.io/",
    "http://localtest.me/",
]


@check
def ssrf_rechazado_por_la_api():
    for url in INTERNAL_URLS:
        r = convert("html-a-pdf", url=url)
        assert r.status_code == 400, f"{url}: {r.status_code} {r.text[:200]}"


GOTENBERG_PROBE = r"""
import json, httpx
urls = json.loads(input())
out = {}
for url in urls:
    files = {"url": (None, url)}
    try:
        r = httpx.post("http://gotenberg:3000/forms/chromium/convert/url", files=files, timeout=60)
        out[url] = r.status_code
    except Exception as exc:
        out[url] = repr(exc)
print(json.dumps(out))
"""


@check
def ssrf_rechazado_por_gotenberg():
    """Sin pasar por la API: Gotenberg mismo (--chromium-deny-private-ips) es la protección de verdad."""
    urls = [
        "http://api:8000/v1/health",
        "http://gotenberg:3000/health",
        "http://127.0.0.1:3000/health",
        "http://169.254.169.254/latest/meta-data/",
        "http://[::1]:3000/health",
        "http://127.0.0.1.nip.io:3000/health",
        "http://localtest.me:3000/health",
    ]
    proc = subprocess.run(
        [*COMPOSE, "exec", "-T", "api", "python", "-c", GOTENBERG_PROBE],
        input=json.dumps(urls),
        capture_output=True,
        text=True,
        cwd=SERVER_DIR,
        timeout=600,
    )
    assert proc.returncode == 0, proc.stderr[-1000:]
    statuses = json.loads(proc.stdout.strip().splitlines()[-1])
    print("   Gotenberg respondió:", statuses)
    for url, status in statuses.items():
        assert status != 200, f"Gotenberg abrió {url}"


PROBE_HTML = b"""<!doctype html><meta charset=utf-8>
<h1>Prueba de seguridad</h1>
<iframe src="file:///tmp/" width=700 height=300></iframe>
<iframe src="../" width=700 height=300></iframe>
<iframe src="../../" width=700 height=300></iframe>
<iframe src="file:///etc/passwd" width=700 height=300></iframe>
<iframe src="http://api:8000/v1/health" width=700 height=100></iframe>
<iframe src="http://gotenberg:3000/health" width=700 height=100></iframe>
<iframe src="http://169.254.169.254/latest/meta-data/" width=700 height=100></iframe>
<img src="http://api:8000/v1/health">
<script>
fetch("http://gotenberg:3000/health").then(r => r.text()).then(t => document.body.append("FETCH:" + t)).catch(() => {});
</script>
"""
FORBIDDEN = ["Index of", "root:x:", "maxUploadMb", '"status"', "ami-id", "FETCH:", "documento.docx", "index.html"]


@check
def html_no_lee_archivos_ni_red_interna():
    """Un HTML subido no puede listar /tmp (archivos de otras conversiones) ni abrir la red interna."""
    from docx import Document

    # Otra conversión al mismo tiempo, para que haya archivos ajenos en /tmp de Gotenberg.
    doc = Document()
    for n in range(400):
        doc.add_paragraph(f"Parrafo {n} de otra persona")
    other = io.BytesIO()
    doc.save(other)
    worker = threading.Thread(target=lambda: convert("word-a-pdf", "ajeno.docx", other.getvalue()))
    worker.start()
    time.sleep(0.5)
    r = convert("html-a-pdf", "sonda.html", PROBE_HTML)
    worker.join()
    content = expect_pdf_with(r, "Prueba de seguridad")
    leaks = [word for word in FORBIDDEN if word in content]
    assert not leaks, f"el PDF muestra {leaks}: {content[:500]!r}"


# ---------- Límites antes de recibir el archivo ----------

STALL = b'--XyZ\r\nContent-Disposition: form-data; name="file"; filename="a.pdf"\r\nContent-Type: application/pdf\r\n\r\n%PDF-1.7\n'


def open_conn(base: str):
    url = urlparse(base)
    port = url.port or (443 if url.scheme == "https" else 80)
    sock = socket.create_connection((url.hostname, port), timeout=15)
    if url.scheme != "https":
        return sock
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE  # certificado propio de Caddy para localhost
    return ctx.wrap_socket(sock, server_hostname=url.hostname)


def request_head(length: int, xff: str | None = None) -> bytes:
    head = f"POST /v1/convert/pdf-a-word HTTP/1.1\r\nHost: localhost\r\nContent-Type: multipart/form-data; boundary=XyZ\r\nContent-Length: {length}\r\n"
    if xff:
        head += f"X-Forwarded-For: {xff}\r\n"
    return head.encode() + b"\r\n"


def upload_and_read(base: str, length: int, xff: str | None = None, seconds: float = 20) -> tuple[int, bytes]:
    """Sube sin parar (como un navegador) y lee la respuesta. Devuelve (bytes enviados, respuesta)."""
    sock = open_conn(base)
    sock.sendall(request_head(length, xff) + STALL)
    sock.setblocking(False)
    sent, response, chunk = 0, b"", b"0" * 16384
    end = time.monotonic() + seconds
    while time.monotonic() < end and sent < length - 100_000:
        try:
            sent += sock.send(chunk)
        except (ssl.SSLWantWriteError, BlockingIOError):
            pass
        except OSError:
            break
        try:
            data = sock.recv(65536)
            if not data:
                break
            response += data
            if b"\r\n\r\n" in response:
                break
        except (ssl.SSLWantReadError, BlockingIOError):
            pass
        except OSError:
            break
        time.sleep(0.001)
    sock.close()
    return sent, response


def status(response: bytes) -> int:
    assert response.startswith(b"HTTP/1.1 "), f"sin respuesta: {response[:100]!r}"
    return int(response.split(b" ", 2)[1])


@check
def archivo_muy_grande_rechazado_al_tiro():
    for base, length in ((API, 60_000_000), (CADDY, 53_000_000), (CADDY, 60_000_000)):
        sent, response = upload_and_read(base, length)
        assert status(response) == 413, f"{base} {length}: {response[:200]!r}"
        assert sent < 15_000_000, f"{base}: el servidor recibió {sent} bytes antes de rechazar"
        print(f"   {base} Content-Length={length}: 413 después de {sent // 1024} KB")


@check
def limite_por_persona():
    # Dos subidas que se quedan a medias ocupan los 2 cupos de esta IP…
    stalled = []
    for _ in range(2):
        sock = open_conn(API)
        sock.sendall(request_head(10_000_000) + STALL)
        stalled.append(sock)
    time.sleep(1)
    # …y la tercera recibe 429 sin que el servidor espere su archivo.
    _, response = upload_and_read(API, 10_000_000)
    assert status(response) == 429, response[:200]
    for sock in stalled:
        sock.close()
    time.sleep(2)


@check
def limite_por_persona_detras_de_caddy():
    """Caddy borra el X-Forwarded-For falso: cambiarlo no sirve para saltarse el límite."""
    stalled = []
    for n in range(2):
        sock = open_conn(CADDY)
        sock.sendall(request_head(10_000_000, f"203.0.113.{n + 1}") + STALL)
        stalled.append(sock)
    time.sleep(1)
    _, response = upload_and_read(CADDY, 10_000_000, "203.0.113.9")
    assert status(response) == 429, response[:200]
    for sock in stalled:
        sock.close()
    time.sleep(2)


@check
def sigue_sano_al_final():
    assert httpx.get(f"{API}/v1/health", timeout=10).json()["ok"] is True
    r = convert("html-a-pdf", "pagina.html", b"<h1>Todavia funciona</h1>")
    expect_pdf_with(r, "Todavia funciona")


if __name__ == "__main__":
    # Sin argumentos corre todo; con nombres, solo esas pruebas (por ejemplo: python integracion.py salud).
    sys.exit(run(sys.argv[1:]))
